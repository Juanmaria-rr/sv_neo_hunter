#!/usr/bin/env python
"""
export_universe_xlsx.py — the candidate universe as one workbook, explained in place.

WHY A SECOND FORMAT
-------------------
The TSV stays canonical: every tool reads it and it diffs cleanly. A TSV cannot
carry its own documentation, though, and a reader who opens the table without
the dictionary beside it reads 130 column names cold. This writes a derived
workbook with the data and its explanation together:

  README             what the file is, its source checksum and the code version
  candidate_universe the table, header frozen, autofilter on; hovering a header
                     shows that column's definition
  column_dictionary  one row per column (candidate_universe_column_dictionary.tsv)
  glossary           the terms the dictionary uses

It refuses to write if any column lacks a definition, or if two columns share
one verbatim — a dictionary that explains several columns with the same
sentence explains none of them.

The workbook is DERIVED. Never edit it; regenerate it from the TSV.

Usage
-----
    python tools/export_universe_xlsx.py --universe-dir <dir> [--out <file.xlsx>]
"""
from __future__ import annotations

import argparse
import hashlib
import pathlib
import re
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from provenance import code_version                     # noqa: E402

FONT = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", fgColor="ECEAE4")
WRAP = Alignment(wrap_text=True, vertical="top")
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")   # rejected by the xlsx format


def sha256(path: pathlib.Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def plain(text) -> str:
    """Dictionary markdown (`code`, **bold**) as plain text for a cell."""
    s = "" if text is None or (isinstance(text, float) and text != text) else str(text)
    return s.replace("**", "").replace("`", "")


def cell_value(v):
    if v is None or (isinstance(v, float) and v != v):
        return None
    if isinstance(v, str):
        return ILLEGAL.sub("", v)
    return v.item() if hasattr(v, "item") else v


def header_row(ws, names, comments=None):
    row = []
    for name in names:
        c = WriteOnlyCell(ws, value=name)
        c.font, c.fill = HEAD, HEAD_FILL
        if comments and comments.get(name):
            c.comment = Comment(comments[name][:2000], "dictionary", width=420, height=220)
        row.append(c)
    ws.append(row)


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__.split("\n")[1])
    p.add_argument("--universe-dir", required=True, type=pathlib.Path)
    p.add_argument("--out", type=pathlib.Path, default=None,
                   help="default: candidate_universe.xlsx beside the TSV")
    args = p.parse_args()

    tsv = args.universe_dir / "candidate_universe.tsv"
    dpath = args.universe_dir / "candidate_universe_column_dictionary.tsv"
    gpath = args.universe_dir / "candidate_universe_column_dictionary_GLOSSARY.tsv"
    out = args.out or args.universe_dir / "candidate_universe.xlsx"

    table = pd.read_csv(tsv, sep="\t", low_memory=False)
    dictionary = pd.read_csv(dpath, sep="\t")
    glossary = pd.read_csv(gpath, sep="\t") if gpath.exists() else pd.DataFrame()

    # ---- coverage: every column explained, and explained on its own
    defined = dictionary.dropna(subset=["what_it_is"])
    defined = defined[defined["what_it_is"].str.strip() != ""]
    missing = [c for c in table.columns if c not in set(defined["column"])]
    extra = [c for c in dictionary["column"] if c not in set(table.columns)]
    shared = defined[defined["what_it_is"].duplicated(keep=False)]["column"].tolist()
    if missing or extra or shared:
        sys.exit(f"REFUSING: dictionary does not match the table.\n"
                 f"  undefined: {missing}\n  not in table: {extra}\n"
                 f"  sharing a definition verbatim: {shared}\n"
                 f"Regenerate it with tools/build_universe_dictionary.py.")

    meaning = {r.column: plain(r.what_it_is) +
               (f"\n\n{plain(r.also_noted)}" if isinstance(r.also_noted, str) and r.also_noted else "")
               for r in dictionary.itertuples()}

    wb = Workbook(write_only=True)

    # ---- README
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 110
    header_row(ws, ["key", "value"])
    readme = [
        ("what this is", "The candidate universe: one row per candidate neopeptide, every "
                         "line of the lineage, nothing filtered out. Filtering is done on "
                         "its columns, never by removing rows."),
        ("derived file", "Generated from the TSV by tools/export_universe_xlsx.py. Do not "
                         "edit; regenerate. The TSV is canonical."),
        ("sheets", "candidate_universe (the data; hover a header for its definition) · "
                   "column_dictionary (every column, its stage, whether it filters, "
                   "how it is derived) · glossary (terms used in the dictionary)"),
        ("empty cells", "An empty cell means the question was not asked (no cohort match, "
                        "no RNA BAM, no gnomAD record), never 'no' or 0."),
        ("counting basis", "acquired_in = the line whose own calls generate the peptide; "
                           "present_in = every line carrying it, inherited included. Filter "
                           "on acquired_in to ask what a line added."),
        ("source", str(tsv)),
        ("source sha256", sha256(tsv)),
        ("dictionary", str(dpath)),
        ("rows × columns", f"{len(table):,} × {table.shape[1]}"),
        ("columns defined", f"{table.shape[1]} of {table.shape[1]}, each with its own text"),
    ] + [(r["key"], r["value"]) for r in code_version()]
    for k, v in readme:
        a, b = WriteOnlyCell(ws, value=k), WriteOnlyCell(ws, value=v)
        a.font, b.font, b.alignment = HEAD, FONT, WRAP
        ws.append([a, b])

    # ---- data
    ws = wb.create_sheet("candidate_universe")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{_col(table.shape[1])}{len(table) + 1}"
    for i, name in enumerate(table.columns, start=1):
        ws.column_dimensions[_col(i)].width = min(max(len(name) + 2, 10), 28)
    header_row(ws, list(table.columns), meaning)
    for values in table.itertuples(index=False, name=None):
        ws.append([cell_value(v) for v in values])

    # ---- dictionary
    ws = wb.create_sheet("column_dictionary")
    ws.freeze_panes = "B2"
    cols = list(dictionary.columns)
    widths = {"column": 30, "stage": 16, "level": 20, "kind": 22, "derived_from": 34,
              "what_it_is": 90, "also_noted": 60, "example": 24}
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[_col(i)].width = widths.get(name, 12)
    header_row(ws, cols)
    for r in dictionary.itertuples(index=False, name=None):
        row = []
        for name, v in zip(cols, r):
            v = plain(v) if name in ("what_it_is", "also_noted", "derived_from") else cell_value(v)
            c = WriteOnlyCell(ws, value=v)
            c.font = HEAD if name == "column" else FONT
            c.alignment = WRAP
            row.append(c)
        ws.append(row)

    # ---- glossary
    if len(glossary):
        ws = wb.create_sheet("glossary")
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 110
        header_row(ws, list(glossary.columns))
        for term, text in glossary.itertuples(index=False, name=None):
            a, b = WriteOnlyCell(ws, value=term), WriteOnlyCell(ws, value=plain(text))
            a.font, b.font, b.alignment = HEAD, FONT, WRAP
            ws.append([a, b])

    wb.save(out)
    print(f"wrote {out}: {len(table):,} rows × {table.shape[1]} columns, "
          f"{len(dictionary)} defined, {len(glossary)} glossary terms", file=sys.stderr)


def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


if __name__ == "__main__":
    main()
